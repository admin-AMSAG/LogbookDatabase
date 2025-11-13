from __future__ import annotations

import csv
import io
import os
from pathlib import Path
from typing import Iterable
from urllib.parse import parse_qs, urlparse

import requests
from flask import (
    Flask,
    Response,
    flash,
    redirect,
    render_template,
    request,
    url_for,
)
from werkzeug.datastructures import FileStorage
from werkzeug.exceptions import BadRequest
import sqlite3

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - handled by requirements
    raise SystemExit("openpyxl must be installed to run this application") from exc

BASE_DIR = Path(__file__).parent
DATABASE_PATH = BASE_DIR / "logbooks.db"
UPLOAD_FOLDER = BASE_DIR / "uploads"
ALLOWED_EXTENSIONS = {"csv", "xlsx", "xls"}
DEFAULT_SHEET_URL = os.environ.get("GOOGLE_SHEET_URL", "").strip()


class ReverseProxied:
    """Allow the app to run behind a reverse proxy or URL prefix.

    When the Flask app is hosted at a sub-path (for example, ``/logbookdatabase``),
    the embedding site or reverse proxy should either set the ``LOGBOOK_APP_PREFIX``
    environment variable or supply the ``X-Script-Name`` header. This middleware
    updates ``SCRIPT_NAME``/``PATH_INFO`` so ``url_for`` continues to generate
    correct links and static asset paths inside an iframe.
    """

    def __init__(self, app, *, script_name: str = "") -> None:
        self.app = app
        self.script_name = script_name.rstrip("/")

    def __call__(self, environ, start_response):
        script_name = environ.get("HTTP_X_SCRIPT_NAME") or self.script_name
        if script_name:
            script_name = script_name.rstrip("/") or ""
            if script_name and not script_name.startswith("/"):
                script_name = f"/{script_name}"
            environ["SCRIPT_NAME"] = script_name
            path_info = environ.get("PATH_INFO", "")
            if script_name and path_info.startswith(script_name):
                environ["PATH_INFO"] = path_info[len(script_name) :] or "/"

        scheme = environ.get("HTTP_X_FORWARDED_PROTO")
        if scheme:
            environ["wsgi.url_scheme"] = scheme

        host = environ.get("HTTP_X_FORWARDED_HOST")
        if host:
            environ["HTTP_HOST"] = host

        return self.app(environ, start_response)


app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("FLASK_SECRET_KEY", "amsag-logbook-secret")
app.config["UPLOAD_FOLDER"] = str(UPLOAD_FOLDER)
app.wsgi_app = ReverseProxied(
    app.wsgi_app, script_name=os.environ.get("LOGBOOK_APP_PREFIX", "")
)
def init_db() -> None:
    """Initialize the database if it does not already exist."""
    with sqlite3.connect(DATABASE_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS logbooks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                logbook_number TEXT UNIQUE NOT NULL,
                owner_name TEXT NOT NULL,
                rego TEXT NOT NULL,
                vehicle TEXT,
                notes TEXT,
                status TEXT NOT NULL DEFAULT 'active'
            )
            """
        )
        conn.commit()


def get_connection() -> sqlite3.Connection:
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


@app.before_first_request
def setup() -> None:
    UPLOAD_FOLDER.mkdir(exist_ok=True)
    init_db()


@app.context_processor
def inject_counts():
    with get_connection() as conn:
        active = conn.execute(
            "SELECT COUNT(*) FROM logbooks WHERE status = 'active'"
        ).fetchone()[0]
        cancelled = conn.execute(
            "SELECT COUNT(*) FROM logbooks WHERE status = 'cancelled'"
        ).fetchone()[0]
    return {"active_count": active, "cancelled_count": cancelled}


@app.route("/")
def index():
    query = request.args.get("query", "").strip()
    filter_status = request.args.get("status", "active")

    sql = "SELECT * FROM logbooks"
    params: list[str] = []
    clauses: list[str] = []

    if query:
        like_query = f"%{query}%"
        clauses.append(
            "(owner_name LIKE ? OR logbook_number LIKE ? OR rego LIKE ?)"
        )
        params.extend([like_query, like_query, like_query])

    if filter_status in {"active", "cancelled"}:
        clauses.append("status = ?")
        params.append(filter_status)

    if clauses:
        sql += " WHERE " + " AND ".join(clauses)

    sql += " ORDER BY owner_name COLLATE NOCASE"

    with get_connection() as conn:
        logbooks = conn.execute(sql, params).fetchall()

    return render_template(
        "index.html", logbooks=logbooks, query=query, filter_status=filter_status
    )


@app.route("/logbooks", methods=["POST"])
def create_logbook():
    form = request.form
    logbook_number = form.get("logbook_number", "").strip()
    owner_name = form.get("owner_name", "").strip()
    rego = form.get("rego", "").strip()
    vehicle = form.get("vehicle", "").strip() or None
    notes = form.get("notes", "").strip() or None

    if not (logbook_number and owner_name and rego):
        flash("Logbook number, owner name, and registration are required.", "error")
        return redirect(url_for("index"))

    with get_connection() as conn:
        try:
            conn.execute(
                """
                INSERT INTO logbooks (logbook_number, owner_name, rego, vehicle, notes)
                VALUES (?, ?, ?, ?, ?)
                """,
                (logbook_number, owner_name, rego, vehicle, notes),
            )
            conn.commit()
            flash("Logbook created successfully.", "success")
        except sqlite3.IntegrityError:
            flash("Logbook number already exists.", "error")

    return redirect(url_for("index"))


@app.route("/logbooks/<int:logbook_id>", methods=["POST"])
def update_logbook(logbook_id: int):
    form = request.form
    owner_name = form.get("owner_name", "").strip()
    rego = form.get("rego", "").strip()
    vehicle = form.get("vehicle", "").strip() or None
    notes = form.get("notes", "").strip() or None
    status = form.get("status", "active").strip()

    if not (owner_name and rego):
        flash("Owner name and registration are required.", "error")
        return redirect(url_for("index"))

    with get_connection() as conn:
        conn.execute(
            """
            UPDATE logbooks
            SET owner_name = ?, rego = ?, vehicle = ?, notes = ?, status = ?
            WHERE id = ?
            """,
            (owner_name, rego, vehicle, notes, status, logbook_id),
        )
        conn.commit()

    flash("Logbook updated successfully.", "success")
    return redirect(url_for("index", status=status))


@app.route("/logbooks/<int:logbook_id>/cancel", methods=["POST"])
def cancel_logbook(logbook_id: int):
    with get_connection() as conn:
        conn.execute(
            "UPDATE logbooks SET status = 'cancelled' WHERE id = ?",
            (logbook_id,),
        )
        conn.commit()
    flash("Logbook marked as cancelled.", "info")
    return redirect(url_for("index", status="cancelled"))


@app.route("/logbooks/<int:logbook_id>/activate", methods=["POST"])
def activate_logbook(logbook_id: int):
    with get_connection() as conn:
        conn.execute(
            "UPDATE logbooks SET status = 'active' WHERE id = ?",
            (logbook_id,),
        )
        conn.commit()
    flash("Logbook reactivated.", "success")
    return redirect(url_for("index"))


@app.route("/import", methods=["GET", "POST"])
def import_data():
    if request.method == "GET":
        return render_template("import.html", default_sheet_url=DEFAULT_SHEET_URL)

    file: FileStorage | None = request.files.get("file")
    sheet_url = request.form.get("sheet_url", "").strip()

    if request.form.get("use_default") and DEFAULT_SHEET_URL:
        sheet_url = DEFAULT_SHEET_URL

    records: list[dict[str, str]] = []

    try:
        if file and file.filename:
            records = parse_upload(file)
        elif sheet_url:
            records = parse_google_sheet(sheet_url)
        else:
            raise BadRequest("Please provide a file or Google Sheet URL.")
    except Exception as exc:  # pragma: no cover - user facing error handling
        flash(f"Import failed: {exc}", "error")
        return redirect(url_for("import_data"))

    imported = upsert_records(records)
    flash(f"Imported {imported} logbooks successfully.", "success")
    return redirect(url_for("index"))


def parse_upload(file: FileStorage) -> list[dict[str, str]]:
    filename = file.filename or ""
    if not allowed_file(filename):
        raise BadRequest("Unsupported file type. Please upload CSV or Excel files.")

    suffix = filename.rsplit(".", 1)[1].lower()

    if suffix == "csv":
        text = file.stream.read().decode("utf-8-sig")
        return list(parse_csv(io.StringIO(text)))

    temp_path = UPLOAD_FOLDER / filename
    file.save(temp_path)
    try:
        return list(parse_excel(temp_path))
    finally:
        temp_path.unlink(missing_ok=True)


def parse_csv(stream: io.TextIOBase) -> Iterable[dict[str, str]]:
    reader = csv.DictReader(stream)
    for row in reader:
        yield normalize_record(row)


def parse_excel(path: Path) -> Iterable[dict[str, str]]:
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(max_row=1))]
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_dict = {headers[idx].lower(): (str(value).strip() if value is not None else "") for idx, value in enumerate(row)}
        yield normalize_record(row_dict)


def normalize_record(row: dict[str, str]) -> dict[str, str]:
    return {
        "logbook_number": row.get("logbook_number", "").strip(),
        "owner_name": row.get("owner_name", "").strip(),
        "rego": row.get("rego", "").strip(),
        "vehicle": row.get("vehicle", "").strip(),
        "notes": row.get("notes", "").strip(),
        "status": row.get("status", "").strip().lower() or "active",
    }


def parse_google_sheet(url: str) -> list[dict[str, str]]:
    csv_url = convert_sheet_url(url)
    response = requests.get(csv_url, timeout=30)
    response.raise_for_status()
    text = response.content.decode("utf-8-sig")
    return list(parse_csv(io.StringIO(text)))


def convert_sheet_url(url: str) -> str:
    if "export?format=csv" in url:
        return url

    parsed = urlparse(url)
    if "docs.google.com" not in parsed.netloc:
        raise BadRequest("Only Google Sheet URLs are supported.")

    path_parts = [part for part in parsed.path.split("/") if part]
    try:
        sheet_index = path_parts.index("d")
        sheet_id = path_parts[sheet_index + 1]
    except (ValueError, IndexError) as exc:
        raise BadRequest("Unrecognized Google Sheet URL.") from exc

    query = parse_qs(parsed.query)
    gid = query.get("gid", [None])[0]
    base = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv"
    if gid:
        return f"{base}&gid={gid}"
    return base


def upsert_records(records: Iterable[dict[str, str]]) -> int:
    inserted = 0
    with get_connection() as conn:
        for record in records:
            if not (record["logbook_number"] and record["owner_name"] and record["rego"]):
                continue
            conn.execute(
                """
                INSERT INTO logbooks (logbook_number, owner_name, rego, vehicle, notes, status)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(logbook_number) DO UPDATE SET
                    owner_name=excluded.owner_name,
                    rego=excluded.rego,
                    vehicle=excluded.vehicle,
                    notes=excluded.notes,
                    status=excluded.status
                """,
                (
                    record["logbook_number"],
                    record["owner_name"],
                    record["rego"],
                    record.get("vehicle") or None,
                    record.get("notes") or None,
                    record.get("status") or "active",
                ),
            )
            inserted += 1
        conn.commit()
    return inserted


@app.route("/export")
def export_data() -> Response:
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT logbook_number, owner_name, rego, vehicle, notes, status FROM logbooks"
        ).fetchall()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["logbook_number", "owner_name", "rego", "vehicle", "notes", "status"])
    for row in rows:
        writer.writerow([row["logbook_number"], row["owner_name"], row["rego"], row["vehicle"] or "", row["notes"] or "", row["status"]])

    output.seek(0)
    return Response(
        output.getvalue(),
        headers={
            "Content-Disposition": "attachment; filename=logbooks.csv",
            "Content-Type": "text/csv",
        },
    )


if __name__ == "__main__":
    init_db()
    app.run(debug=True)
